import openpyxl
import json
import re
import os

excel_path = r'e:\x01-git\xDashBoard\-doc\UHT-InstruementList.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb['Equipment Tags (2)']

headers = [ws.cell(4, c).value for c in range(1, 20)]

parsed_json_items = []
raw_rows = []

for r in range(5, ws.max_row + 1):
    col1 = ws.cell(r, 1).value
    if col1 is None:
        continue
    row_vals = [ws.cell(r, c).value for c in range(1, 20)]
    row_dict = {}
    for h, v in zip(headers, row_vals):
        row_dict[h] = v if v is not None else ""
    
    # Clean string values
    for k in row_dict:
        if isinstance(row_dict[k], str):
            row_dict[k] = row_dict[k].strip()

    raw_rows.append(row_dict)

    # For parsed_instruments.json
    parsed_json_items.append({
        "Column1": row_dict.get("Column1", ""),
        "Process": row_dict.get("Process", ""),
        "PLC": row_dict.get("PLC", "192.168.0.101"),
        "Database": row_dict.get("Database", ""),
        "Label": row_dict.get("Label", ""),
        "DashBoard": row_dict.get("DashBoard", ""),
        "ID": row_dict.get("ID", ""),
        "Group": row_dict.get("Group", ""),
        "Type": row_dict.get("Type3", ""),
        "Signal Type": row_dict.get("Signal Type", ""),
        "Min": str(row_dict.get("Min", "")),
        "Max": str(row_dict.get("Max", "")),
        "Setpoint": str(row_dict.get("Setpoint", "")),
        "UOM": row_dict.get("UOM", ""),
        "Description": row_dict.get("Description", ""),
        "Denomination": row_dict.get("Denomination", ""),
        "Type2": row_dict.get("Type2", ""),
        "Doc No.": row_dict.get("Doc No.", ""),
        "Note": row_dict.get("Note", ""),
        "Rev.": ""
    })

# Save parsed_instruments.json
json_out_path = r'e:\x01-git\xDashBoard\parsed_instruments.json'
with open(json_out_path, 'w', encoding='utf-8') as f:
    json.dump(parsed_json_items, f, indent=2, ensure_ascii=False)

print(f"Generated {json_out_path} with {len(parsed_json_items)} items.")

# Now build mysql/init.sql
# 1. Collect unique groups
groups_set = []
group_map = {} # group_code -> group_id
group_id_counter = 1

# Default groups from raw_rows
for item in raw_rows:
    g_code = item.get("Group", "").strip()
    if not g_code:
        g_code = "General"
    if g_code not in group_map:
        group_map[g_code] = group_id_counter
        groups_set.append((group_id_counter, g_code, f"{g_code} Subsystem"))
        group_id_counter += 1

# 2. Collect unique hardware catalog items
hardware_list = []
hardware_map = {} # (denom, spec, doc) -> hw_id
hw_id_counter = 1

for item in raw_rows:
    denom = item.get("Denomination", "").strip()
    spec = item.get("Type2", "").strip()
    doc = item.get("Doc No.", "").strip()

    if denom or spec or doc:
        key = (denom, spec, doc)
        if key not in hardware_map:
            hardware_map[key] = hw_id_counter
            hardware_list.append((hw_id_counter, denom, spec, doc))
            hw_id_counter += 1

# 3. Parse process_instruments
process_instruments = []
alarm_configs = []

inst_id_counter = 1

for item in raw_rows:
    tag_name = item.get("Label", "").strip()
    if not tag_name:
        tag_name = f"TAG_{inst_id_counter}"
    
    # Database string parsing e.g. DB1601.DBD0000, DB1601.DBX100.0, DB1601.DBW960, DB1600.DBD0176
    db_str = item.get("Database", "").strip()
    m = re.match(r'DB(\d+)\.(DBD|DBW|DBX)(\d+)(?:\.(\d+))?', db_str, re.IGNORECASE)
    if m:
        db_num = int(m.group(1))
        db_type = m.group(2).upper()
        byte_off = float(int(m.group(3)))
        bit_off = int(m.group(4)) if m.group(4) else 0
    else:
        db_num = 1601
        db_type = 'DBD'
        byte_off = 0.0
        bit_off = 0

    type3 = item.get("Type3", "").strip()
    if type3.lower() == 'real':
        plc_format = 'REAL'
    elif type3.lower() in ('int', 'word'):
        plc_format = 'INT'
    elif type3.lower() in ('bool', 'boolean'):
        plc_format = 'BOOL'
    else:
        plc_format = 'REAL'

    sig_type_raw = item.get("Signal Type", "").strip()
    if 'on/off' in sig_type_raw.lower():
        sig_type = 'ON_OFF'
    elif 'hz' in sig_type_raw.lower() or 'current' in sig_type_raw.lower() or 'power' in sig_type_raw.lower():
        sig_type = 'ENG_VALUE'
    elif 'pid' in sig_type_raw.lower():
        sig_type = 'ENG_VALUE'
    elif 'status' in sig_type_raw.lower():
        sig_type = 'ON_OFF'
    else:
        sig_type = 'ENG_VALUE'

    g_code = item.get("Group", "").strip() or "General"
    g_id = group_map[g_code]

    denom = item.get("Denomination", "").strip()
    spec = item.get("Type2", "").strip()
    doc = item.get("Doc No.", "").strip()
    hw_id = hardware_map.get((denom, spec, doc), None)

    dashboard_flag = 'TRUE' if item.get("DashBoard", "").strip().upper() == 'Y' else 'FALSE'

    try:
        scale_min = float(item.get("Min", 0) or 0)
    except:
        scale_min = 0.0
    
    try:
        scale_max = float(item.get("Max", 100) or 100)
    except:
        scale_max = 100.0

    setpoint_val = item.get("Setpoint", "")
    try:
        setpoint = f"{float(setpoint_val):.2f}" if setpoint_val != "" and setpoint_val is not None else "NULL"
    except:
        setpoint = "NULL"

    uom = item.get("UOM", "").strip()
    desc = item.get("Description", "").strip().replace("'", "''")

    scale_raw_min = 0.0
    scale_raw_max = 27648.0 if plc_format == 'REAL' else 1.0

    process_instruments.append({
        "id": inst_id_counter,
        "area_id": 1,
        "plc_id": 1,
        "group_id": g_id,
        "hardware_id": hw_id if hw_id else "NULL",
        "tag_name": tag_name,
        "description": desc,
        "show_on_dashboard": dashboard_flag,
        "db_number": db_num,
        "byte_offset": f"{byte_off:.1f}",
        "bit_offset": bit_off,
        "plc_data_format": plc_format,
        "signal_type": sig_type,
        "scale_raw_min": f"{scale_raw_min:.1f}",
        "scale_raw_max": f"{scale_raw_max:.1f}",
        "scale_eng_min": f"{scale_min:.2f}",
        "scale_eng_max": f"{scale_max:.2f}",
        "target_setpoint": setpoint,
        "engineering_unit": uom
    })

    # Check for alarm configuration (e.g. HH, H, L, LL tags or setpoint)
    if 'HH' in tag_name:
        alarm_configs.append((inst_id_counter, 'HIGH_HIGH', setpoint if setpoint != "NULL" else "40.0"))
    elif 'H' in tag_name and 'HH' not in tag_name and 'SH' not in tag_name:
        alarm_configs.append((inst_id_counter, 'HIGH', setpoint if setpoint != "NULL" else "30.0"))
    elif 'LL' in tag_name:
        alarm_configs.append((inst_id_counter, 'LOW_LOW', setpoint if setpoint != "NULL" else "10.0"))
    elif 'L' in tag_name and 'LL' not in tag_name and 'SL' not in tag_name:
        alarm_configs.append((inst_id_counter, 'LOW', setpoint if setpoint != "NULL" else "20.0"))
    elif sig_type == 'ON_OFF' and setpoint != "NULL":
        alarm_configs.append((inst_id_counter, 'DIGITAL_SWITCH', setpoint))

    inst_id_counter += 1

# Format SQL output
sql_lines = [
    "-- ==========================================================================",
    "-- SIEMENS S7 PLC DATABLOCK-AWARE SCADA DATABASE INITIALIZATION",
    "-- Source File: UHT-InstruementList.xlsx (Updated 80 Instruments)",
    "-- Target PLC IP: 192.168.0.101",
    "-- ==========================================================================",
    "",
    "CREATE DATABASE IF NOT EXISTS thpack_dashboard DEFAULT CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;",
    "USE thpack_dashboard;",
    "",
    "SET FOREIGN_KEY_CHECKS = 0;",
    "DROP VIEW IF EXISTS vw_siemens_plc_db_map;",
    "DROP VIEW IF EXISTS vw_scada_live_metrics;",
    "DROP VIEW IF EXISTS vw_instrument_alarms;",
    "DROP VIEW IF EXISTS vw_instruments_full_spec;",
    "DROP VIEW IF EXISTS vw_scada_dashboard_tags;",
    "DROP TABLE IF EXISTS telemetry_readings;",
    "DROP TABLE IF EXISTS alarm_configurations;",
    "DROP TABLE IF EXISTS process_instruments;",
    "DROP TABLE IF EXISTS sensor_hardware_catalog;",
    "DROP TABLE IF EXISTS instrument_groups;",
    "DROP TABLE IF EXISTS plc_gateways;",
    "DROP TABLE IF EXISTS process_areas;",
    "SET FOREIGN_KEY_CHECKS = 1;",
    "",
    "CREATE TABLE process_areas (",
    "    id INT AUTO_INCREMENT PRIMARY KEY,",
    "    area_code VARCHAR(50) NOT NULL UNIQUE,",
    "    area_name VARCHAR(100) NOT NULL,",
    "    description TEXT,",
    "    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP",
    ") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;",
    "",
    "CREATE TABLE plc_gateways (",
    "    id INT AUTO_INCREMENT PRIMARY KEY,",
    "    area_id INT NOT NULL,",
    "    gateway_name VARCHAR(100) NOT NULL,",
    "    ip_address VARCHAR(50) NOT NULL DEFAULT '192.168.0.101',",
    "    rack INT NOT NULL DEFAULT 0,",
    "    slot INT NOT NULL DEFAULT 1,",
    "    poll_interval_ms INT NOT NULL DEFAULT 1000,",
    "    status ENUM('CONNECTED', 'DISCONNECTED', 'ERROR', 'SIMULATED') DEFAULT 'SIMULATED',",
    "    last_connected_at TIMESTAMP NULL,",
    "    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,",
    "    FOREIGN KEY (area_id) REFERENCES process_areas(id) ON DELETE CASCADE",
    ") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;",
    "",
    "CREATE TABLE instrument_groups (",
    "    id INT AUTO_INCREMENT PRIMARY KEY,",
    "    area_id INT NOT NULL,",
    "    group_code VARCHAR(50) NOT NULL,",
    "    group_name VARCHAR(100) NOT NULL,",
    "    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,",
    "    UNIQUE KEY uk_area_group (area_id, group_code),",
    "    FOREIGN KEY (area_id) REFERENCES process_areas(id) ON DELETE CASCADE",
    ") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;",
    "",
    "CREATE TABLE sensor_hardware_catalog (",
    "    id INT AUTO_INCREMENT PRIMARY KEY,",
    "    denomination VARCHAR(255) NOT NULL,",
    "    model_specification TEXT,",
    "    doc_number VARCHAR(100),",
    "    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP",
    ") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;",
    "",
    "CREATE TABLE process_instruments (",
    "    id INT AUTO_INCREMENT PRIMARY KEY,",
    "    area_id INT NOT NULL,",
    "    plc_id INT NOT NULL,",
    "    group_id INT NOT NULL,",
    "    hardware_id INT,",
    "    tag_name VARCHAR(100) NOT NULL,",
    "    description TEXT,",
    "    show_on_dashboard BOOLEAN DEFAULT TRUE,",
    "    db_number INT NOT NULL DEFAULT 1601,",
    "    byte_offset DECIMAL(8,2) NOT NULL DEFAULT 0.00,",
    "    bit_offset INT NOT NULL DEFAULT 0,",
    "    plc_data_format ENUM('REAL', 'INT', 'BOOL', 'WORD', 'DWORD') NOT NULL DEFAULT 'REAL',",
    "    signal_type ENUM('ENG_VALUE', 'ON_OFF', 'ALARM', 'CONTROL') NOT NULL DEFAULT 'ENG_VALUE',",
    "    scale_raw_min DECIMAL(10,2) DEFAULT 0.00,",
    "    scale_raw_max DECIMAL(10,2) DEFAULT 27648.00,",
    "    scale_eng_min DECIMAL(10,2) DEFAULT 0.00,",
    "    scale_eng_max DECIMAL(10,2) DEFAULT 100.00,",
    "    target_setpoint DECIMAL(10,2),",
    "    engineering_unit VARCHAR(50),",
    "    document_revision VARCHAR(50),",
    "    remarks TEXT,",
    "    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,",
    "    FOREIGN KEY (area_id) REFERENCES process_areas(id) ON DELETE CASCADE,",
    "    FOREIGN KEY (plc_id) REFERENCES plc_gateways(id) ON DELETE CASCADE,",
    "    FOREIGN KEY (group_id) REFERENCES instrument_groups(id) ON DELETE CASCADE,",
    "    FOREIGN KEY (hardware_id) REFERENCES sensor_hardware_catalog(id) ON DELETE SET NULL",
    ") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;",
    "",
    "CREATE TABLE alarm_configurations (",
    "    id INT AUTO_INCREMENT PRIMARY KEY,",
    "    instrument_id INT NOT NULL,",
    "    alarm_type ENUM('HIGH_HIGH', 'HIGH', 'LOW', 'LOW_LOW', 'DIGITAL_SWITCH') NOT NULL,",
    "    threshold_value DECIMAL(10,2),",
    "    is_enabled BOOLEAN DEFAULT TRUE,",
    "    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,",
    "    FOREIGN KEY (instrument_id) REFERENCES process_instruments(id) ON DELETE CASCADE",
    ") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;",
    "",
    "CREATE TABLE telemetry_readings (",
    "    id BIGINT AUTO_INCREMENT PRIMARY KEY,",
    "    instrument_id INT NOT NULL,",
    "    value_scaled DECIMAL(12,4),",
    "    value_raw DECIMAL(12,4),",
    "    value_digital BOOLEAN,",
    "    quality_status ENUM('GOOD', 'BAD', 'UNCERTAIN') DEFAULT 'GOOD',",
    "    recorded_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,",
    "    FOREIGN KEY (instrument_id) REFERENCES process_instruments(id) ON DELETE CASCADE,",
    "    INDEX idx_inst_time (instrument_id, recorded_at)",
    ") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;",
    "",
    "INSERT INTO process_areas (id, area_code, area_name, description) VALUES (1, 'UHT-01', 'UHT Processing Area UHT-01', 'Aseptic Processing Plant Area UHT-01');",
    "INSERT INTO plc_gateways (id, area_id, gateway_name, ip_address, rack, slot, poll_interval_ms, status) VALUES (1, 1, 'Siemens S7-1500 UHT PLC', '192.168.0.101', 0, 1, 1000, 'SIMULATED');",
    ""
]

for gid, gcode, gdesc in groups_set:
    gcode_esc = gcode.replace("'", "''")
    gdesc_esc = gdesc.replace("'", "''")
    sql_lines.append(f"INSERT INTO instrument_groups (id, area_id, group_code, group_name) VALUES ({gid}, 1, '{gcode_esc}', '{gdesc_esc}');")

sql_lines.append("")

for hwid, denom, spec, doc in hardware_list:
    denom_esc = denom.replace("'", "''")
    spec_esc = f"'{spec.replace('\'', '\'\'')}'" if spec else "NULL"
    doc_esc = f"'{doc.replace('\'', '\'\'')}'" if doc else "NULL"
    sql_lines.append(f"INSERT INTO sensor_hardware_catalog (id, denomination, model_specification, doc_number) VALUES ({hwid}, '{denom_esc}', {spec_esc}, {doc_esc});")

sql_lines.append("")

for pi in process_instruments:
    tag = pi['tag_name'].replace("'", "''")
    uom_esc = f"'{pi['engineering_unit'].replace('\'', '\'\'')}'" if pi['engineering_unit'] else "NULL"
    sql_lines.append(
        f"INSERT INTO process_instruments (id, area_id, plc_id, group_id, hardware_id, tag_name, description, show_on_dashboard, db_number, byte_offset, bit_offset, plc_data_format, signal_type, scale_raw_min, scale_raw_max, scale_eng_min, scale_eng_max, target_setpoint, engineering_unit) "
        f"VALUES ({pi['id']}, {pi['area_id']}, {pi['plc_id']}, {pi['group_id']}, {pi['hardware_id']}, '{tag}', '{pi['description']}', {pi['show_on_dashboard']}, {pi['db_number']}, {pi['byte_offset']}, {pi['bit_offset']}, '{pi['plc_data_format']}', '{pi['signal_type']}', {pi['scale_raw_min']}, {pi['scale_raw_max']}, {pi['scale_eng_min']}, {pi['scale_eng_max']}, {pi['target_setpoint']}, {uom_esc});"
    )

sql_lines.append("")

for aid, (inst_id, atype, thresh) in enumerate(alarm_configs, start=1):
    sql_lines.append(f"INSERT INTO alarm_configurations (id, instrument_id, alarm_type, threshold_value, is_enabled) VALUES ({aid}, {inst_id}, '{atype}', {thresh}, TRUE);")

sql_lines.extend([
    "",
    "-- Analytical Views",
    "CREATE VIEW vw_siemens_plc_db_map AS",
    "SELECT ",
    "    i.id AS instrument_id,",
    "    g.gateway_name,",
    "    g.ip_address,",
    "    i.tag_name,",
    "    i.db_number,",
    "    i.byte_offset,",
    "    i.bit_offset,",
    "    i.plc_data_format,",
    "    i.scale_raw_min,",
    "    i.scale_raw_max,",
    "    i.scale_eng_min,",
    "    i.scale_eng_max,",
    "    i.engineering_unit,",
    "    i.description",
    "FROM process_instruments i",
    "JOIN plc_gateways g ON i.plc_id = g.id",
    "ORDER BY i.db_number, i.byte_offset, i.bit_offset;",
    "",
    "CREATE VIEW vw_scada_live_metrics AS",
    "SELECT ",
    "    i.id AS instrument_id,",
    "    a.area_code,",
    "    ig.group_code,",
    "    i.tag_name,",
    "    i.description,",
    "    i.plc_data_format,",
    "    i.signal_type,",
    "    i.scale_eng_min,",
    "    i.scale_eng_max,",
    "    i.target_setpoint,",
    "    i.engineering_unit,",
    "    h.denomination AS hardware_type",
    "FROM process_instruments i",
    "JOIN process_areas a ON i.area_id = a.id",
    "JOIN instrument_groups ig ON i.group_id = ig.id",
    "LEFT JOIN sensor_hardware_catalog h ON i.hardware_id = h.id",
    "WHERE i.show_on_dashboard = TRUE;",
    "",
    "CREATE VIEW vw_instrument_alarms AS",
    "SELECT ",
    "    ac.id AS alarm_id,",
    "    i.tag_name,",
    "    a.area_code,",
    "    i.description,",
    "    ac.alarm_type,",
    "    ac.threshold_value,",
    "    i.engineering_unit,",
    "    ac.is_enabled",
    "FROM alarm_configurations ac",
    "JOIN process_instruments i ON ac.instrument_id = i.id",
    "JOIN process_areas a ON i.area_id = a.id;",
    ""
])

sql_out_path = r'e:\x01-git\xDashBoard\mysql\init.sql'
with open(sql_out_path, 'w', encoding='utf-8') as f:
    f.write("\n".join(sql_lines))

print(f"Generated {sql_out_path} with {len(sql_lines)} lines.")
